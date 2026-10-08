import base64
import csv
import io
import logging
from datetime import date, datetime

from odoo import _, fields, models
from odoo.exceptions import UserError, ValidationError

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:
    openpyxl = None


def _decode_csv_content(content):
    """Helper to decode CSV bytes supporting UTF-8 with BOM and Latin-1 fallback."""
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("latin-1")


def _parse_date_value(val):
    """Helper to parse a date string or object across standard formats."""
    if not val:
        return None
    if isinstance(val, (datetime, date)):
        return val.date() if isinstance(val, datetime) else val
    val_str = str(val).strip()
    if not val_str:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(val_str, fmt).date()
        except ValueError:
            pass
    return None


class WizardImportPublicHoliday(models.TransientModel):
    _name = "wizard.import.public.holiday"
    _description = "Import Public Holidays Wizard"

    file_data = fields.Binary(string="Holiday File (.csv, .xlsx)", required=False)
    file_name = fields.Char(string="File Name")
    target_year = fields.Integer(
        string="Target Year",
        default=lambda self: fields.Date.today().year,
        help="Year for generating standard Cambodian fixed holidays.",
    )
    duplicate_policy = fields.Selection(
        [
            ("skip", "Skip Existing"),
            ("update", "Update Existing"),
        ],
        string="Duplicate Policy",
        default="skip",
        required=True,
        help="How to handle holidays that already exist with the same name and start date.",
    )
    state = fields.Selection(
        [
            ("upload", "Upload"),
            ("preview", "Preview & Validate"),
            ("done", "Completed"),
        ],
        string="State",
        default="upload",
    )
    preview_line_ids = fields.One2many(
        "wizard.import.public.holiday.line",
        "wizard_id",
        string="Preview Lines",
    )
    total_rows = fields.Integer(string="Total Rows", default=0)
    valid_rows = fields.Integer(string="Valid Rows", default=0)
    duplicate_rows = fields.Integer(string="Duplicate Rows", default=0)
    error_rows = fields.Integer(string="Error Rows", default=0)
    summary_msg = fields.Text(string="Summary")
    import_log_id = fields.Many2one("school.holiday.import.log", string="Import Log")

    def action_download_template(self):
        """Generate and return sample CSV template."""
        csv_content = (
            "Name,Name KM,Date From,Date To,Holiday Type,Work Pay Multiplier,Note\n"
            "International New Year,ទិវាចូលឆ្នាំសកល,2026-01-01,2026-01-01,fixed,2.0,International New Year\n"
            "Khmer New Year,ពិធីបុណ្យចូលឆ្នាំថ្មីប្រពៃណីជាតិ,2026-04-14,2026-04-16,fixed,2.0,National Khmer New Year Celebration\n"
            "King's Birthday,ព្រះរាជពិធីបុណ្យចម្រើនព្រះជន្ម,2026-05-14,2026-05-14,fixed,2.0,King Norodom Sihamoni Birthday\n"
            "Pchum Ben,ពិធីបុណ្យភ្ជុំបិណ្ឌ,2026-10-10,2026-10-12,lunar,2.0,Ancestors Day Festival\n"
        )
        attachment = self.env["ir.attachment"].create({
            "name": "public_holidays_sample_template.csv",
            "type": "binary",
            "datas": base64.b64encode(csv_content.encode("utf-8")),
            "mimetype": "text/csv",
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{attachment.id}?download=true",
            "target": "self",
        }

    def action_generate_fixed_holidays(self):
        """Generate standard Cambodian fixed-date public holidays for the specified target year."""
        self.ensure_one()
        year = self.target_year or fields.Date.today().year
        fixed_definitions = [
            ("International New Year", "ទិវាចូលឆ្នាំសកល", f"{year}-01-01", f"{year}-01-01"),
            ("Victory over Genocide Day", "ទិវាជ័យជម្នះលើរបបប្រល័យពូជសាសន៍", f"{year}-01-07", f"{year}-01-07"),
            ("International Women's Day", "ទិវាអន្តរជាតិនារី", f"{year}-03-08", f"{year}-03-08"),
            ("Khmer New Year Day 1", "ពិធីបុណ្យចូលឆ្នាំថ្មីប្រពៃណីជាតិ ថ្ងៃទី១", f"{year}-04-14", f"{year}-04-14"),
            ("Khmer New Year Day 2", "ពិធីបុណ្យចូលឆ្នាំថ្មីប្រពៃណីជាតិ ថ្ងៃទី២", f"{year}-04-15", f"{year}-04-15"),
            ("Khmer New Year Day 3", "ពិធីបុណ្យចូលឆ្នាំថ្មីប្រពៃណីជាតិ ថ្ងៃទី៣", f"{year}-04-16", f"{year}-04-16"),
            ("International Labour Day", "ទិវាពលកម្មអន្តរជាតិ", f"{year}-05-01", f"{year}-05-01"),
            ("King's Birthday", "ព្រះរាជពិធីបុណ្យចម្រើនព្រះជន្មព្រះមហាក្សត្រ", f"{year}-05-14", f"{year}-05-14"),
            ("Constitution Day", "ទិវាប្រកាសរដ្ឋធម្មនុញ្ញ", f"{year}-09-24", f"{year}-09-24"),
            ("King Father's Memorial Day", "ទិវាគោរពព្រះវិញ្ញាណក្ខន្ធព្រះបរមរតនកោដ្ឋ", f"{year}-10-15", f"{year}-10-15"),
            ("Coronation Day", "ព្រះរាជពិធីគ្រងរាជសម្បត្តិ", f"{year}-10-29", f"{year}-10-29"),
            ("Independence Day", "ទិវាបុណ្យឯករាជ្យជាតិ", f"{year}-11-09", f"{year}-11-09"),
        ]
        PublicHoliday = self.env["public.holiday"]
        created_count = 0
        for name, name_km, d_from_s, d_to_s in fixed_definitions:
            d_from = fields.Date.to_date(d_from_s)
            d_to = fields.Date.to_date(d_to_s)
            existing = PublicHoliday.search([
                ("name", "=ilike", name),
                ("date_from", "=", d_from),
                ("active", "=", True),
            ], limit=1)
            if not existing:
                PublicHoliday.create({
                    "name": name,
                    "name_km": name_km,
                    "date_from": d_from,
                    "date_to": d_to,
                    "holiday_type": "fixed",
                    "affects_classes": True,
                    "active": True,
                })
                created_count += 1
        return {
            "type": "ir.actions.client",
            "tag": "display_notification",
            "params": {
                "title": _("Fixed Holidays Generated"),
                "message": _("Generated %d fixed public holidays for year %d.") % (created_count, year),
                "type": "success",
                "sticky": False,
            },
        }

    def _extract_rows_from_file(self):
        self.ensure_one()
        if not self.file_data:
            raise UserError(_("Please upload a CSV or Excel file."))

        content = base64.b64decode(self.file_data)
        file_name = (self.file_name or "").lower()
        rows = []

        if file_name.endswith(".xlsx") or file_name.endswith(".xls"):
            if not openpyxl:
                raise UserError(_("openpyxl library is required to read Excel files."))
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            sheet = wb.active
            iter_rows = sheet.iter_rows(values_only=True)
            try:
                headers = [str(cell).strip().lower() if cell is not None else "" for cell in next(iter_rows)]
            except StopIteration:
                return []
            for row in iter_rows:
                if not any(row):
                    continue
                row_dict = {}
                for idx, cell in enumerate(row):
                    if idx < len(headers):
                        row_dict[headers[idx]] = cell
                rows.append(row_dict)
        else:
            text = _decode_csv_content(content)
            reader = csv.reader(io.StringIO(text))
            try:
                raw_headers = next(reader)
                headers = [h.strip().lower() for h in raw_headers]
            except StopIteration:
                return []
            for row in reader:
                if not row or not any(row):
                    continue
                row_dict = {}
                for idx, cell in enumerate(row):
                    if idx < len(headers):
                        row_dict[headers[idx]] = cell.strip()
                rows.append(row_dict)

        return rows

    def _get_val(self, row, aliases, default=None):
        for alias in aliases:
            if alias in row and row[alias] is not None and str(row[alias]).strip() != "":
                return row[alias]
        return default

    def action_parse_and_preview(self):
        self.ensure_one()
        raw_rows = self._extract_rows_from_file()
        if not raw_rows:
            raise UserError(_("The uploaded file contains no data rows."))

        Holiday = self.env["public.holiday"]
        preview_vals = []
        valid_cnt = 0
        duplicate_cnt = 0
        error_cnt = 0

        type_mapping = {
            "fixed": "fixed",
            "public": "fixed",
            "public holiday": "fixed",
            "lunar": "lunar",
            "special": "special",
            "school_break": "special",
            "school break": "special",
            "exam_break": "special",
            "exam break": "special",
        }

        seen_keys = set()
        for idx, row in enumerate(raw_rows, start=2):
            name = self._get_val(row, ["name", "holiday", "holiday name", "title"])
            name_km = self._get_val(row, ["name_km", "name km", "khmer name", "khmer"])
            d_from_raw = self._get_val(row, ["date from", "date_from", "start date", "date_start", "start", "from"])
            d_to_raw = self._get_val(row, ["date to", "date_to", "end date", "date_end", "end", "to"])
            h_type_raw = self._get_val(row, ["holiday type", "holiday_type", "type"], "fixed")
            mult_raw = self._get_val(row, ["work pay multiplier", "work_pay_multiplier", "multiplier", "rate"], 2.0)
            note = self._get_val(row, ["note", "notes", "description"], "")

            line_status = "ok"
            status_msg = _("Valid row")

            if not name:
                line_status = "missing_name"
                status_msg = _("Missing holiday name")
                error_cnt += 1
            else:
                name = str(name).strip()

            d_from = _parse_date_value(d_from_raw)
            if not d_from and line_status == "ok":
                line_status = "invalid_date"
                status_msg = _("Invalid Start Date: '%s'") % d_from_raw
                error_cnt += 1

            d_to = _parse_date_value(d_to_raw) if d_to_raw else d_from
            if not d_to and d_to_raw and line_status == "ok":
                line_status = "invalid_date"
                status_msg = _("Invalid End Date: '%s'") % d_to_raw
                error_cnt += 1

            if d_from and d_to and d_to < d_from and line_status == "ok":
                line_status = "invalid_date"
                status_msg = _("End Date (%s) is before Start Date (%s)") % (d_to, d_from)
                error_cnt += 1

            h_type = type_mapping.get(str(h_type_raw).strip().lower(), "fixed")

            try:
                mult = float(mult_raw)
            except (ValueError, TypeError):
                mult = 2.0

            if line_status == "ok" and name and d_from:
                pair_key = (name.lower(), d_from)
                if pair_key in seen_keys:
                    line_status = "duplicate"
                    status_msg = _("Duplicate holiday in file: %s on %s") % (name, d_from)
                    duplicate_cnt += 1
                else:
                    seen_keys.add(pair_key)
                    existing = Holiday.search([
                        ("name", "=ilike", name),
                        ("date_from", "=", d_from),
                        ("active", "=", True),
                    ], limit=1)
                    if existing:
                        line_status = "duplicate"
                        status_msg = _("Matches existing holiday ID %d in database") % existing.id
                        duplicate_cnt += 1
                    else:
                        valid_cnt += 1

            preview_vals.append((0, 0, {
                "line_number": idx,
                "name": name or "",
                "name_km": name_km or "",
                "date_from": d_from,
                "date_to": d_to,
                "holiday_type": h_type,
                "work_pay_multiplier": mult,
                "note": str(note or ""),
                "status": line_status,
                "status_message": status_msg,
            }))

        self.preview_line_ids.unlink()
        self.write({
            "preview_line_ids": preview_vals,
            "total_rows": len(raw_rows),
            "valid_rows": valid_cnt,
            "duplicate_rows": duplicate_cnt,
            "error_rows": error_cnt,
            "state": "preview",
        })

        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    def action_confirm_import(self):
        self.ensure_one()
        if not self.preview_line_ids:
            raise UserError(_("No preview lines found. Please upload and preview first."))

        fatal_lines = self.preview_line_ids.filtered(lambda l: l.status in ("missing_name", "invalid_date"))
        if fatal_lines and len(fatal_lines) == len(self.preview_line_ids):
            raise ValidationError(
                _("Cannot proceed with import: %d rows contain invalid dates or missing names. Please fix your file and retry.")
                % len(fatal_lines)
            )

        PubHoliday = self.env["public.holiday"]
        UnivHoliday = self.env["university.holiday"]
        created_records = PubHoliday
        updated_count = 0
        skipped_count = 0
        failed_count = 0
        error_msgs = []

        log = self.env["school.holiday.import.log"].create({
            "name": self.file_name or _("Public Holidays Import"),
            "duplicate_policy": self.duplicate_policy,
            "total_records": len(self.preview_line_ids),
        })

        for line in self.preview_line_ids:
            if line.status in ("missing_name", "invalid_date"):
                failed_count += 1
                error_msgs.append(f"Row {line.line_number}: {line.status_message}")
                continue

            try:
                with self.env.cr.savepoint():
                    if line.status == "duplicate":
                        existing_pub = PubHoliday.search([
                            ("name", "=ilike", line.name),
                            ("date_from", "=", line.date_from),
                        ], limit=1)
                        if self.duplicate_policy == "update":
                            if existing_pub:
                                existing_pub.write({
                                    "date_to": line.date_to or line.date_from,
                                    "name_km": line.name_km or existing_pub.name_km,
                                    "holiday_type": line.holiday_type if line.holiday_type in ("fixed", "lunar", "special") else "fixed",
                                    "note": line.note,
                                })
                            existing_univ = UnivHoliday.search([
                                ("name", "=ilike", line.name),
                                ("date_from", "=", line.date_from),
                            ], limit=1)
                            if existing_univ:
                                existing_univ.write({
                                    "date_to": line.date_to or line.date_from,
                                    "date_end": line.date_to or line.date_from,
                                    "work_pay_multiplier": line.work_pay_multiplier,
                                    "note": line.note,
                                })
                            updated_count += 1
                        else:
                            skipped_count += 1
                        continue

                    created = PubHoliday.create({
                        "name": line.name,
                        "name_km": line.name_km or False,
                        "date_from": line.date_from,
                        "date_to": line.date_to or line.date_from,
                        "holiday_type": line.holiday_type if line.holiday_type in ("fixed", "lunar", "special") else "fixed",
                        "affects_classes": True,
                        "active": True,
                        "note": line.note,
                    })
                    created_records |= created

                    # Also create institutional closure record if university.holiday exists
                    existing_univ = UnivHoliday.search([
                        ("name", "=ilike", line.name),
                        ("date_from", "=", line.date_from),
                    ], limit=1)
                    if not existing_univ:
                        UnivHoliday.create({
                            "name": line.name,
                            "date_from": line.date_from,
                            "date_to": line.date_to or line.date_from,
                            "date_start": line.date_from,
                            "date_end": line.date_to or line.date_from,
                            "holiday_type": "public",
                            "applies_to": "all",
                            "work_pay_multiplier": line.work_pay_multiplier,
                            "note": line.note,
                            "import_batch_id": log.id,
                        })
            except Exception as exc:
                failed_count += 1
                error_msgs.append(f"Row {line.line_number} ({line.name}): {str(exc)}")

        log.write({
            "records_created": len(created_records),
            "records_updated": updated_count,
            "records_skipped": skipped_count,
            "records_failed": failed_count,
            "error_log": "\n".join(error_msgs) if error_msgs else _("Completed cleanly without errors."),
        })

        self.write({
            "state": "done",
            "import_log_id": log.id,
            "summary_msg": _(
                "Import Completed Successfully!\n"
                "• Created: %d\n"
                "• Updated: %d\n"
                "• Skipped: %d\n"
                "• Failed: %d\n"
                "• Total Processed: %d"
            )
            % (len(created_records), updated_count, skipped_count, failed_count, len(self.preview_line_ids)),
        })

        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }

    def action_check_class_conflicts(self):
        self.ensure_one()
        action = self.env.ref("school_management.action_university_schedule_conflicts", raise_if_not_found=False)
        if action:
            return action.read()[0]
        return {
            "type": "ir.actions.act_window",
            "name": _("Schedule Conflicts"),
            "res_model": "university.timetable.conflict",
            "view_mode": "kanban,list,form",
        }

    def action_view_imported_holidays(self):
        self.ensure_one()
        if self.import_log_id:
            return self.import_log_id.action_view_holidays()
        return {"type": "ir.actions.act_window_close"}


class WizardImportPublicHolidayLine(models.TransientModel):
    _name = "wizard.import.public.holiday.line"
    _description = "Holiday Import Preview Line"
    _order = "line_number asc"

    wizard_id = fields.Many2one("wizard.import.public.holiday", ondelete="cascade", required=True)
    line_number = fields.Integer(string="Line #")
    name = fields.Char(string="Holiday Name")
    name_km = fields.Char(string="Khmer Name")
    date_from = fields.Date(string="Start Date")
    date_to = fields.Date(string="End Date")
    holiday_type = fields.Selection(
        [
            ("fixed", "Fixed"),
            ("lunar", "Lunar"),
            ("special", "Special"),
        ],
        string="Type",
        default="fixed",
    )
    work_pay_multiplier = fields.Float(string="Multiplier", default=2.0)
    note = fields.Text(string="Note")
    status = fields.Selection(
        [
            ("ok", "Valid"),
            ("duplicate", "Duplicate"),
            ("missing_name", "Missing Name"),
            ("invalid_date", "Invalid Date"),
        ],
        string="Status",
        default="ok",
    )
    status_message = fields.Char(string="Validation Message")
